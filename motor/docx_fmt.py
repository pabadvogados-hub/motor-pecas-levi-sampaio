"""Gera o .docx a partir dos blocos da minuta (python-docx)."""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.shared import Cm, Pt

LACUNA = re.compile(r"(\[[^\[\]]*●\])")
ALINH = {
    "justificado": WD_ALIGN_PARAGRAPH.JUSTIFY,
    "esquerda": WD_ALIGN_PARAGRAPH.LEFT,
}


def _runs(par, texto: str, fonte: str, tamanho: int, negrito=False, italico=False):
    """Escreve o texto destacando em amarelo cada lacuna [... ●]."""
    for pedaco in LACUNA.split(texto):
        if not pedaco:
            continue
        r = par.add_run(pedaco)
        r.font.name = fonte
        r.font.size = Pt(tamanho)
        r.bold = negrito
        r.italic = italico
        if "●" in pedaco:
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW


def gerar(blocos: list[dict], destino: Path, formato: dict, timbrado: Path | None = None) -> Path:
    usa_timbrado = bool(timbrado and Path(timbrado).exists())
    doc = Document(str(timbrado)) if usa_timbrado else Document()
    if usa_timbrado:  # remove corpo de exemplo, preserva cabecalho/rodape/margens
        for el in list(doc.element.body):
            if not el.tag.endswith("sectPr"):
                doc.element.body.remove(el)
    else:
        m = formato["margens_cm"]
        for s in doc.sections:
            s.top_margin, s.bottom_margin = Cm(m["topo"]), Cm(m["base"])
            s.left_margin, s.right_margin = Cm(m["esquerda"]), Cm(m["direita"])

    fonte, tam = formato["fonte"], formato["tamanho_corpo"]
    alinh = ALINH.get(formato.get("alinhamento", "justificado"), WD_ALIGN_PARAGRAPH.JUSTIFY)

    for b in blocos:
        k, t = b["k"], b["t"]
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.line_spacing = formato["espacamento"]
        pf.space_after = Pt(6)
        if k == "endereco":
            p.alignment = ALINH["justificado"]
            _runs(p, t.upper(), fonte, tam, negrito=True)
            pf.space_after = Pt(24)
        elif k == "processo":
            p.alignment = ALINH["esquerda"]
            _runs(p, t, fonte, tam, negrito=True)
            pf.space_after = Pt(18)
        elif k == "titulo":
            p.alignment = alinh
            pf.space_before = Pt(12)
            pf.keep_with_next = True
            _runs(p, t.upper(), fonte, tam, negrito=True)
        elif k == "subtitulo":
            p.alignment = alinh
            pf.keep_with_next = True
            _runs(p, t, fonte, tam, negrito=True)
        elif k == "nota":
            p.alignment = ALINH["esquerda"]
            r = p.add_run(t)
            r.font.name, r.font.size, r.italic = fonte, Pt(tam - 1), True
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
        elif k == "item":
            p.alignment = alinh
            pf.left_indent = Cm(1.25)
            _runs(p, t, fonte, tam)
        elif k == "centro":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _runs(p, t, fonte, tam)
        elif k == "citacao":
            p.alignment = alinh
            pf.left_indent = Cm(4)
            pf.line_spacing = 1.0
            _runs(p, t, fonte, tam - 1, italico=True)
        else:  # paragrafo comum
            p.alignment = alinh
            pf.first_line_indent = Cm(formato["recuo_primeira_linha_cm"])
            _runs(p, t, fonte, tam)

    destino.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(destino))
    return destino


def blocos_de_markdown(md: str) -> list[dict]:
    """Converte o texto devolvido pelo Claude (markdown simples) em blocos."""
    blocos: list[dict] = []
    for linha in md.splitlines():
        s = linha.strip()
        if not s:
            continue
        if s.startswith("# "):
            blocos.append({"k": "endereco", "t": s[2:]})
        elif s.startswith("## "):
            blocos.append({"k": "titulo", "t": s[3:]})
        elif s.startswith("### "):
            blocos.append({"k": "subtitulo", "t": s[4:]})
        elif s.startswith("> "):
            blocos.append({"k": "citacao", "t": s[2:]})
        elif re.match(r"^[a-z]\)\s", s):
            blocos.append({"k": "item", "t": s})
        elif s.startswith("[REVISAR"):
            blocos.append({"k": "nota", "t": s})
        else:
            blocos.append({"k": "p", "t": s.replace("**", "")})
    return blocos
